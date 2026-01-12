from pathlib import Path
import pandas as pd

def extract_text_from_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except Exception as e:
        raise RuntimeError(
            "pdfplumber is required for PDF extraction. "
            "Install with: pip install pdfplumber"
        ) from e

    text_parts = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                text_parts.append(text)
    return "\n\n".join(text_parts)


def extract_text_from_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as e:
        raise RuntimeError(
            "python-docx is required for DOCX extraction. "
            "Install with: pip install python-docx"
        ) from e

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def extract_text_from_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_text_from_parquet(path: Path, nrows: int = 1000) -> str:
    df = pd.read_parquet(path)
    rows = []
    for _, r in df.head(nrows).iterrows():
        row_text = " | ".join([f"{c}: {str(r[c])}" for c in df.columns])
        rows.append(row_text)
    return "\n".join(rows)


def extract_text_from_pdf_bytes(content: bytes) -> str:
    text = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            if page.extract_text():
                text.append(page.extract_text())
    return "\n".join(text)


def extract_text_from_docx_bytes(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs)


def extract_text_from_parquet(path: Path) -> str:
    df = pd.read_parquet(path)
    return df.astype(str).to_csv(index=False)
